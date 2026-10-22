"""Tests for the Excel → catalog write-back in core.excel_io.

A planner exports the database workbook, edits a few labor minutes and adds a
SKU row on the ``Machine catalog`` / ``Accessory catalog`` sheet, and uploads
it back. The reader has to find those sheets and reconcile their header with
the live CSV; the diff has to report exactly the cells that would change —
numbers compared as numbers, blanks never applied, absent SKUs untouched, and
anything that doesn't validate blocking the save rather than slipping through.
"""
from __future__ import annotations

import ast
import datetime as dt
import io
import re
from pathlib import Path

import pandas as pd
import pytest

from core import excel_io
from core.excel_io import (
    CatalogImportError,
    diff_catalog_sheet,
    normalize_routing_value,
    parse_labor_value,
    read_catalog_workbook,
)

APP = Path(__file__).resolve().parent.parent / "app.py"

MACHINE_COLS = [
    "SKU", "Description", "Warehouse", "Wire", "Trailer", "FN_Assy", "PDI",
    "QC", "Ship", "ETO", "Bat", "Last Modified",
    "Final Station HND", "Final Station CYP", "Wire Station HND",
]
ACC_COLS = [
    "SKU", "Description", "Warehouse", "AccKIT", "Nameplate Prep", "BattSubRaw",
    "PMAcc", "GenAcc", "Last Modified", "ComAcc",
]
LIVE = {"machine": MACHINE_COLS, "accessory": ACC_COLS}


def _workbook(sheets: dict[str, list[list]]) -> bytes:
    """Build an .xlsx in memory from {sheet name: raw rows} — rows are written
    exactly as given so tests control where the header sits."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, index=False, header=False)
    return buf.getvalue()


def _machine_rows(*rows):
    """``[header, *rows]`` padded to the machine header's width."""
    out = [MACHINE_COLS]
    for r in rows:
        out.append(list(r) + [None] * (len(MACHINE_COLS) - len(r)))
    return out


def _machine_live() -> pd.DataFrame:
    """What load_machine_labor returns for three SKUs: numbers as floats, blank
    routing cells already defaulted to the station, index = SKU."""
    df = pd.DataFrame([
        {"SKU": "BOSS25-010", "Description": "EBOSS 25 TRAILER", "Warehouse": 10.0, "Wire": 60.0,
         "Trailer": 210.0, "FN_Assy": 45.0, "PDI": 60.0, "QC": 60.0, "Ship": 30.0, "ETO": 0.0,
         "Bat": 1, "Last Modified": "2026-05-26",
         "Final Station HND": "Final", "Final Station CYP": "Final", "Wire Station HND": "Wire"},
        {"SKU": "BOSS70-002", "Description": "EBOSS 70", "Warehouse": 10.0, "Wire": 90.0,
         "Trailer": 240.0, "FN_Assy": 0.0, "PDI": 60.0, "QC": 60.0, "Ship": 30.0, "ETO": 0.0,
         "Bat": 3, "Last Modified": "",
         "Final Station HND": "Final", "Final Station CYP": "Final", "Wire Station HND": "Wire"},
        {"SKU": "PDS185-001", "Description": "PDS 185", "Warehouse": 5.0, "Wire": 0.0,
         "Trailer": 0.0, "FN_Assy": 120.0, "PDI": 30.0, "QC": 30.0, "Ship": 15.0, "ETO": 0.0,
         "Bat": 0, "Last Modified": "",
         "Final Station HND": "Undercarriage", "Final Station CYP": "Undercarriage",
         "Wire Station HND": "Wire"},
    ])
    return df.set_index("SKU", drop=False)


def _acc_live() -> pd.DataFrame:
    df = pd.DataFrame([
        {"SKU": "BOSS25-A016", "Description": "KIT A", "Warehouse": 10.0, "AccKIT": 10.0,
         "Nameplate Prep": 0.0, "BattSubRaw": 320.0, "PMAcc": 60.0, "GenAcc": 0.0,
         "Last Modified": "", "ComAcc": 0.0},
    ])
    return df.set_index("SKU", drop=False)


def _sheet(book: bytes, key: str = "machine", live=LIVE):
    return read_catalog_workbook(book, "edited.xlsx", live_columns=live).sheets[key]


# ---------------------------------------------------------------------------
# Reader — finding the sheets
# ---------------------------------------------------------------------------
class TestReadCatalogWorkbook:
    def test_exported_workbook_yields_both_catalogs_and_skips_the_rest(self):
        book = _workbook({
            "README": [["Sheet", "Source file"], ["Machine catalog", "data/machine_clean.csv"]],
            "Machine catalog": _machine_rows(["BOSS25-010", "EBOSS 25 TRAILER", "10.0", "60.0"]),
            "Accessory catalog": [ACC_COLS, ["BOSS25-A016", "KIT A", "10.0", "10.0"]],
            "Crew": [["Facility", "Station", "HC"], ["Henderson", "QC", 4]],
        })
        wb = read_catalog_workbook(book, "db.xlsx", live_columns=LIVE)
        assert set(wb.sheets) == {"machine", "accessory"}
        assert wb.sheets["machine"].matched_by == "name"
        assert wb.sheets["machine"].sheet_name == "Machine catalog"
        assert wb.sheets["machine"].file_path == "data/machine_clean.csv"
        assert wb.sheets["accessory"].file_path == "data/acc_clean.csv"
        assert [n for n, _why in wb.skipped] == ["README", "Crew"]
        assert wb.sheet_names == ["README", "Machine catalog", "Accessory catalog", "Crew"]

    @pytest.mark.parametrize("name", ["machine CATALOG", " Machine catalog ", "machines", "machine_clean"])
    def test_sheet_name_matching_is_forgiving(self, name):
        wb = read_catalog_workbook(_workbook({name: _machine_rows(["X", "d"])}), live_columns=LIVE)
        assert list(wb.sheets) == ["machine"]

    def test_single_unnamed_sheet_is_identified_by_its_header(self):
        book = _workbook({"Sheet1": _machine_rows(["BOSS25-010", "x", "10"])})
        wb = read_catalog_workbook(book, live_columns=LIVE)
        assert list(wb.sheets) == ["machine"]
        assert wb.sheets["machine"].matched_by == "header"
        acc = _workbook({"Sheet1": [ACC_COLS, ["BOSS25-A016", "KIT A", "10"]]})
        assert list(read_catalog_workbook(acc, live_columns=LIVE).sheets) == ["accessory"]

    def test_header_below_a_title_row_is_found(self):
        rows = [["Machine catalog — edited 2026-10-20"], [None]] + _machine_rows(["BOSS25-010", "x", "10"])
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        assert sheet.frame["SKU"].tolist() == ["BOSS25-010"]
        assert sheet.columns[:3] == ["SKU", "Description", "Warehouse"]

    def test_ambiguous_sheet_is_skipped_with_a_reason(self):
        # SKU + Warehouse only: both catalogs have those, so it's no answer.
        book = _workbook({"Sheet1": [["SKU", "Warehouse"], ["X", 1]],
                          "Notes": [["hello"], ["world"]]})
        wb = read_catalog_workbook(book, live_columns=LIVE)
        assert wb.sheets == {}
        reasons = dict(wb.skipped)
        assert "labor columns" in reasons["Sheet1"]
        assert "SKU" in reasons["Notes"]

    def test_named_sheet_without_a_sku_header_is_skipped_not_fatal(self):
        book = _workbook({"Machine catalog": [["Part", "Minutes"], ["X", 1]],
                          "Accessory catalog": [ACC_COLS, ["BOSS25-A016", "KIT A"]]})
        wb = read_catalog_workbook(book, live_columns=LIVE)
        assert list(wb.sheets) == ["accessory"]
        assert dict(wb.skipped)["Machine catalog"].startswith("no header row")

    def test_second_sheet_for_the_same_catalog_is_skipped(self):
        book = _workbook({"Machine catalog": _machine_rows(["A", "x"]),
                          "machines": _machine_rows(["B", "y"])})
        wb = read_catalog_workbook(book, live_columns=LIVE)
        assert wb.sheets["machine"].sheet_name == "Machine catalog"
        assert "already supplies" in dict(wb.skipped)["machines"]

    def test_garbage_bytes_raise_planner_friendly_error(self):
        with pytest.raises(CatalogImportError) as exc:
            read_catalog_workbook(b"definitely not a zip", "bad.xlsx", live_columns=LIVE)
        assert "bad.xlsx" in str(exc.value)

    def test_live_columns_default_to_the_data_dir(self, tmp_path, monkeypatch):
        (tmp_path / "machine_clean.csv").write_text("SKU,Description,Wire,Extra\nX,Y,1,2\n", encoding="utf-8")
        cols = excel_io.live_catalog_columns(tmp_path)
        assert cols["machine"] == ["SKU", "Description", "Wire", "Extra"]
        assert cols["accessory"] == []  # file missing → empty, no crash
        monkeypatch.setattr(excel_io, "DATA_DIR", tmp_path)
        sheet = _sheet(_workbook({"Machine catalog": [["SKU", "Wire", "Trailer"], ["X", 5, 6]]}),
                       live=None)
        assert sheet.columns == ["SKU", "Wire"]
        assert sheet.extra_columns == ["Trailer"]

    def test_real_catalog_headers_are_recognised(self):
        # The deployed CSV headers (what the export writes) classify correctly.
        cols = excel_io.live_catalog_columns()
        book = _workbook({"Sheet1": [cols["machine"]]})
        assert list(read_catalog_workbook(book, live_columns=cols).sheets) == ["machine"]
        book = _workbook({"Sheet1": [cols["accessory"]]})
        assert list(read_catalog_workbook(book, live_columns=cols).sheets) == ["accessory"]


# ---------------------------------------------------------------------------
# Reader — header validation and cell normalisation
# ---------------------------------------------------------------------------
class TestHeaderValidation:
    def test_extra_columns_ignored_and_missing_columns_reported(self):
        header = ["SKU", "Description", "Warehouse", "Notes", "Final Station HND"]
        sheet = _sheet(_workbook({"Machine catalog": [header, ["BOSS25-010", "x", "10", "hi", "Final"]]}))
        assert sheet.columns == ["SKU", "Description", "Warehouse", "Final Station HND"]
        assert sheet.extra_columns == ["Notes"]
        assert "Notes" not in sheet.frame.columns
        # Every editable live column the sheet lacks is listed (SKU / Last
        # Modified are never editable, so they're not).
        assert "Wire" in sheet.missing_columns and "Final Station CYP" in sheet.missing_columns
        assert "SKU" not in sheet.missing_columns and "Last Modified" not in sheet.missing_columns

    def test_header_cells_match_case_insensitively_and_trimmed(self):
        header = ["sku", " Warehouse ", "FN_ASSY", "final station hnd"]
        sheet = _sheet(_workbook({"Machine catalog": [header, ["X", 1, 2, "final"]]}))
        assert sheet.columns == ["SKU", "Warehouse", "FN_Assy", "Final Station HND"]
        assert sheet.extra_columns == []

    def test_duplicate_header_first_column_wins(self):
        sheet = _sheet(_workbook({"Machine catalog": [["SKU", "Wire", "Wire"], ["X", 1, 2]]}))
        assert sheet.columns == ["SKU", "Wire"]
        assert sheet.frame["Wire"].tolist() == ["1"]

    def test_blank_sku_rows_dropped_and_cells_rendered_as_text(self):
        rows = _machine_rows(
            ["BOSS25-010", "x", 10.0, 60, "12.5", None, dt.datetime(2026, 5, 1)],
            [None, "orphan", 1],
            ["", "", ""],
            [" BOSS70-002 ", "y"],
        )
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        assert sheet.frame["SKU"].tolist() == ["BOSS25-010", "BOSS70-002"]
        first = sheet.frame.iloc[0]
        assert first["Warehouse"] == "10"       # whole float → no ".0"
        assert first["Wire"] == "60"
        assert first["Trailer"] == "12.5"
        assert first["FN_Assy"] == ""           # blank cell → "", never "nan"
        assert first["PDI"] == "2026-05-01"     # a date cell is text, caught later as not-a-number

    def test_duplicate_skus_reported_case_insensitively(self):
        rows = _machine_rows(["BOSS25-010", "a"], ["boss25-010", "b"], ["BOSS70-002", "c"])
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        assert sheet.duplicate_skus == ["BOSS25-010"]

    def test_editable_columns_cover_labor_description_and_every_routing_column(self):
        numeric, text = excel_io.editable_columns(excel_io.CATALOG_SPECS["machine"], MACHINE_COLS)
        assert numeric == ["Warehouse", "Wire", "Trailer", "FN_Assy", "PDI", "QC", "Ship", "ETO", "Bat"]
        assert text == ["Description", "Final Station HND", "Final Station CYP", "Wire Station HND"]
        numeric, text = excel_io.editable_columns(excel_io.CATALOG_SPECS["accessory"], ACC_COLS)
        assert numeric == ["Warehouse", "AccKIT", "Nameplate Prep", "BattSubRaw", "PMAcc", "GenAcc", "ComAcc"]
        assert text == ["Description"]

    def test_routing_options_match_the_in_app_editor(self):
        # The import validates routing cells against the same options the
        # machine editor's dropdowns offer. Parse those lists out of app.py so
        # a change to one side without the other fails here.
        source = APP.read_text(encoding="utf-8")
        for station, options in excel_io.ROUTING_OPTIONS.items():
            m = re.search(
                r'SelectboxColumn\(\s*"' + re.escape(station) + r' Station",\s*options=(\[[^\]]*\])',
                source,
            )
            assert m, f"no SelectboxColumn for `{station} Station` in app.py"
            assert tuple(ast.literal_eval(m.group(1))) == options, station


# ---------------------------------------------------------------------------
# Value coercion
# ---------------------------------------------------------------------------
class TestParseLaborValue:
    @pytest.mark.parametrize("raw, expected", [
        ("45", 45.0), ("45.0", 45.0), (45, 45.0), (45.0, 45.0), ("12.5", 12.5),
        (" 1,200 ", 1200.0), ("1e3", 1000.0), ("0", 0.0),
    ])
    def test_numbers(self, raw, expected):
        assert parse_labor_value(raw) == expected

    @pytest.mark.parametrize("raw", ["", None, "   ", float("nan")])
    def test_blank_means_no_value(self, raw):
        assert parse_labor_value(raw) is None

    @pytest.mark.parametrize("raw, why", [
        ("abc", "not a number"), ("4O", "not a number"), ("-5", "negative"),
        ("inf", "finite"), ("2026-05-01", "not a number"),
    ])
    def test_rejections_say_why(self, raw, why):
        with pytest.raises(ValueError) as exc:
            parse_labor_value(raw)
        assert why in str(exc.value)


class TestNormalizeRoutingValue:
    def test_case_is_normalised_to_the_dropdown_spelling(self):
        assert normalize_routing_value("Final Station HND", "final") == "Final"
        assert normalize_routing_value("PDI Station CYP", "qc") == "QC"
        assert normalize_routing_value("Undercarriage Station SPB", "COMACC") == "ComAcc"

    def test_blank_stays_blank(self):
        assert normalize_routing_value("Final Station HND", "") == ""
        assert normalize_routing_value("Final Station HND", None) == ""

    def test_option_from_another_station_is_rejected(self):
        # Wire is a valid station but not somewhere Final-Assembly labor can go.
        with pytest.raises(ValueError) as exc:
            normalize_routing_value("Final Station HND", "Wire")
        assert "not one of" in str(exc.value) and "Undercarriage" in str(exc.value)

    def test_non_routing_column_is_passed_through(self):
        assert normalize_routing_value("Description", " hello ") == "hello"


# ---------------------------------------------------------------------------
# Diff — the preview and what the save applies
# ---------------------------------------------------------------------------
class TestDiffCatalogSheet:
    def test_only_the_edited_cells_change(self):
        rows = _machine_rows(
            # Wire 60 → 75, Description edited, PDI "60.0" and FN_Assy 45 unchanged,
            # Final Station HND → Undercarriage
            ["BOSS25-010", "EBOSS 25 TRAILER (rev B)", "10.0", 75, "210.0", 45, "60.0", "60.0",
             "30.0", "0.0", "1.0", "2026-05-26", "Undercarriage", "Final", "Wire"],
            # untouched row, exported verbatim
            ["BOSS70-002", "EBOSS 70", "10.0", "90.0", "240.0", "0.0", "60.0", "60.0", "30.0",
             "0.0", "3.0", "", "Final", "Final", "Wire"],
        )
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.cell_tuples() == [
            ("BOSS25-010", "Wire", 75.0),
            ("BOSS25-010", "Description", "EBOSS 25 TRAILER (rev B)"),
            ("BOSS25-010", "Final Station HND", "Undercarriage"),
        ]
        assert diff.cells[0].old == 60.0
        assert diff.changed_skus == ["BOSS25-010"]
        assert diff.new_rows == [] and diff.problems == []
        assert not diff.blocked and not diff.is_empty
        assert diff.n_rows == 2 and diff.n_existing == 2
        assert diff.absent_skus == ["PDS185-001"]

    def test_blank_cells_never_apply(self):
        # A planner clears a few cells (or a never-edited row has blank routing
        # in the CSV while the loader shows the default): nothing changes.
        rows = _machine_rows(["BOSS25-010", "", "", None, "", "", "", "", "", "", "", "", "", "", ""])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.is_empty and not diff.blocked

    def test_numbers_compare_as_numbers(self):
        rows = _machine_rows(["BOSS25-010", "EBOSS 25 TRAILER", 10, "60", "210.00", "45.0", 60.0,
                              " 60 ", "30", "0", 1])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.is_empty

    def test_absent_skus_are_untouched_never_deleted(self):
        rows = _machine_rows(["PDS185-001", "PDS 185", "5"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.absent_skus == ["BOSS25-010", "BOSS70-002"]
        assert diff.is_empty  # and no cell for the absent rows, of course

    def test_new_sku_becomes_a_new_row(self):
        rows = _machine_rows(["BOSS25-011", "EBOSS 25 NEW", "10", "60", "", "45", "60", "60", "30",
                              "", "1", "", "final", "", "wire"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.new_skus == ["BOSS25-011"]
        row = diff.new_rows[0]
        assert row["Wire"] == 60.0 and row["Bat"] == 1.0
        assert row["Trailer"] == "" and row["ETO"] == ""          # blank stays blank (loader → 0)
        assert row["Final Station HND"] == "Final"               # normalised spelling
        assert row["Final Station CYP"] == ""                     # blank routing → default at load
        assert row["Description"] == "EBOSS 25 NEW"
        assert "Last Modified" not in row and diff.cells == []
        assert diff.n_existing == 0

    def test_sku_matching_is_case_insensitive_and_keeps_catalog_spelling(self):
        rows = _machine_rows(["boss25-010", "EBOSS 25 TRAILER", "11"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.cell_tuples() == [("BOSS25-010", "Warehouse", 11.0)]
        assert diff.new_rows == []
        assert "BOSS25-010" not in diff.absent_skus

    def test_missing_columns_change_nothing(self):
        # Sheet carries only SKU + Wire: no other column can produce a diff.
        sheet = _sheet(_workbook({"Machine catalog": [["SKU", "Wire"], ["BOSS25-010", 61]]}))
        diff = diff_catalog_sheet(sheet, _machine_live())
        assert diff.cell_tuples() == [("BOSS25-010", "Wire", 61.0)]
        assert diff.numeric_cols == ["Wire"] and diff.text_cols == []

    def test_routing_column_missing_from_a_stale_cache_still_diffs(self):
        live = _machine_live().drop(columns=["Final Station CYP"])
        rows = _machine_rows(["BOSS25-010", "", "", "", "", "", "", "", "", "", "", "", "", "GenAcc"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), live)
        assert diff.cell_tuples() == [("BOSS25-010", "Final Station CYP", "GenAcc")]
        assert diff.cells[0].old == ""

    def test_unparseable_number_is_a_problem_and_blocks(self):
        rows = _machine_rows(
            ["BOSS25-010", "", "", "4O"],                 # letter O for zero
            ["BOSS70-002", "", "", "", "-5"],
            ["NEW-1", "new", "abc"],
        )
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.blocked
        assert [(p.sku, p.column, p.value) for p in diff.problems] == [
            ("BOSS25-010", "Wire", "4O"), ("BOSS70-002", "Trailer", "-5"), ("NEW-1", "Warehouse", "abc"),
        ]
        assert "not a number" in diff.problems[0].reason
        assert "negative" in diff.problems[1].reason
        assert diff.cells == [] and diff.new_rows == []   # the bad new row is not added
        frame = diff.problems_frame()
        assert list(frame.columns) == ["SKU", "Column", "Value", "Problem"] and len(frame) == 3

    def test_invalid_routing_value_is_a_problem(self):
        rows = _machine_rows(["BOSS25-010"] + [""] * 11 + ["Wire"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.blocked
        assert diff.problems[0].column == "Final Station HND"
        assert "not one of" in diff.problems[0].reason

    @pytest.mark.parametrize("sku", ["+CMD()", "-CMD", "@CMD"])
    def test_formula_prefixed_new_sku_is_rejected(self, sku):
        # (A cell starting with "=" can't even be written as text by openpyxl
        # — it becomes a formula — so the other three prefixes stand in for
        # the whole guard here.)
        rows = _machine_rows([sku, "evil", "1"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.blocked and diff.new_rows == []
        assert diff.problems[0].column == "SKU" and "formula" in diff.problems[0].reason

    def test_equals_prefixed_sku_is_rejected_too(self):
        # Straight to the diff with a hand-built sheet, since the workbook
        # writer can't carry a literal "=..." string.
        sheet = _sheet(_workbook({"Machine catalog": _machine_rows(["X", "x", "1"])}))
        sheet.frame.loc[0, "SKU"] = "=CMD()"
        diff = diff_catalog_sheet(sheet, _machine_live())
        assert diff.blocked and diff.problems[0].value == "=CMD()"

    def test_duplicate_skus_block_and_are_left_out(self):
        rows = _machine_rows(["BOSS25-010", "", "11"], ["boss25-010", "", "12"], ["BOSS70-002", "", "13"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        assert diff.blocked and diff.duplicate_skus == ["BOSS25-010"]
        assert diff.cell_tuples() == [("BOSS70-002", "Warehouse", 13.0)]
        # The duplicated SKU is simply excluded — not reported as absent.
        assert "BOSS25-010" not in diff.absent_skus

    def test_accessory_sheet(self):
        rows = [ACC_COLS, ["BOSS25-A016", "KIT A", "10", "10", "0", "340", "60", "0", "", "0"],
                ["BOSS25-A017", "KIT B", "10", "10", "", "320", "", "", "", "118"]]
        diff = diff_catalog_sheet(_sheet(_workbook({"Accessory catalog": rows}), "accessory"), _acc_live())
        assert diff.cell_tuples() == [("BOSS25-A016", "BattSubRaw", 340.0)]
        assert diff.new_skus == ["BOSS25-A017"]
        assert diff.new_rows[0]["ComAcc"] == 118.0 and diff.new_rows[0]["PMAcc"] == ""

    def test_preview_frames_are_display_strings(self):
        rows = _machine_rows(
            ["BOSS25-010", "EBOSS 25 TRAILER", "10", "72.5", "210", "45", "60", "60", "30", "0", "1",
             "", "ComAcc"],
            ["NEW-1", "brand new", "4", "", "", "", "", "", "", "", "2"],
        )
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), _machine_live())
        cells = diff.cells_frame()
        assert list(cells.columns) == ["SKU", "Column", "Old", "New", "Δ"]
        assert cells.iloc[0].tolist() == ["BOSS25-010", "Wire", "60", "72.5", "+12.5"]
        assert cells.iloc[1].tolist() == ["BOSS25-010", "Final Station HND", "Final", "ComAcc", ""]
        new = diff.new_rows_frame()
        assert new.iloc[0]["SKU"] == "NEW-1" and new.iloc[0]["Warehouse"] == "4"
        assert new.iloc[0]["Bat"] == "2" and new.iloc[0]["Wire"] == ""

    def test_empty_live_catalog_makes_everything_new(self):
        rows = _machine_rows(["A", "a", "1"], ["B", "b", "2"])
        diff = diff_catalog_sheet(_sheet(_workbook({"Machine catalog": rows})), pd.DataFrame())
        assert diff.new_skus == ["A", "B"] and diff.absent_skus == []


# ---------------------------------------------------------------------------
# app.py — new rows + the SHA-locked save with precomputed diffs
# ---------------------------------------------------------------------------
app = pytest.importorskip("app")
from core.catalog_storage import GitHubConflict  # noqa: E402


def _fresh_csv_frame() -> pd.DataFrame:
    """The catalog as it is on GitHub at save time. BOSS70-002's Wire was
    changed to 999 by someone else after the planner exported — the import
    must leave that alone."""
    return pd.read_csv(io.StringIO(
        "SKU,Description,Warehouse,Wire,Trailer,Bat,Last Modified,Final Station HND\n"
        "BOSS25-010,EBOSS 25 TRAILER,10.0,60.0,210.0,1.0,2026-05-26,Final\n"
        "BOSS70-002,EBOSS 70,10.0,999.0,240.0,3.0,,Final\n"
        "PDS185-001,PDS 185,5.0,0.0,0.0,0.0,,Undercarriage\n"
    ))


class TestAppendCatalogRows:
    def test_appends_with_blank_fill_and_todays_stamp(self):
        out, added, skipped = app._append_catalog_rows(
            _fresh_csv_frame(),
            [{"SKU": "NEW-1", "Description": "new", "Wire": 12.0, "Trailer": "", "Final Station HND": "Final"}],
        )
        assert added == ["NEW-1"] and skipped == []
        row = out[out["SKU"] == "NEW-1"].iloc[0]
        assert row["Wire"] == 12.0 and row["Description"] == "new"
        assert row["Trailer"] == "" and row["Warehouse"] == "" and row["Bat"] == ""
        assert row["Last Modified"] == app.today_local_str()
        assert len(out) == 4

    def test_skips_existing_formula_prefixed_and_batch_duplicates(self):
        out, added, skipped = app._append_catalog_rows(
            _fresh_csv_frame(),
            [{"SKU": "boss25-010", "Wire": 1},      # already on file (case-insensitive)
             {"SKU": "=HACK", "Wire": 1},
             {"SKU": "NEW-2", "Wire": 1}, {"SKU": "new-2", "Wire": 2},
             {"SKU": "", "Wire": 3}],
        )
        assert added == ["NEW-2"]
        assert skipped == ["boss25-010", "=HACK", "new-2"]
        assert len(out) == 4

    def test_unknown_columns_in_an_addition_are_dropped(self):
        out, added, _ = app._append_catalog_rows(_fresh_csv_frame(), [{"SKU": "N", "Nope": 1}])
        assert added == ["N"] and "Nope" not in out.columns


class TestSaveCatalogCsvWithPrecomputedDiffs:
    @pytest.fixture
    def github(self, monkeypatch):
        """Stub the GitHub round-trip: record every PUT, optionally conflict first."""
        calls = []
        state = {"conflicts": 0}

        def fake_read(file_path, token):
            return _fresh_csv_frame(), "sha-fresh"

        def fake_save(csv_text, file_path, token, message, sha=None):
            if state["conflicts"] > 0:
                state["conflicts"] -= 1
                raise GitHubConflict("409")
            calls.append({"csv": csv_text, "path": file_path, "token": token,
                          "message": message, "sha": sha})
            return {"commit": {"sha": "abc1234def", "html_url": "https://x/commit/abc1234def"}}

        monkeypatch.setattr(app, "_read_fresh_csv", fake_read)
        monkeypatch.setattr(app, "save_catalog_to_github", fake_save)
        monkeypatch.setattr(app, "_track_and_flush", lambda *a, **k: None)
        monkeypatch.setattr(app.st, "secrets", {"github_token": "tok"})
        return calls, state

    def _save(self, diffs, additions):
        return app._save_catalog_csv(
            edited=None, source_df=_machine_live(), editable_cols=list(excel_io.MACHINE_LABOR_COLUMNS),
            file_path="data/machine_clean.csv", label="machine",
            text_cols=["Description", "Final Station HND"],
            diffs=diffs, additions=additions, via="Excel import",
        )

    def test_only_the_diffed_cells_and_new_row_land(self, github):
        calls, _ = github
        result = self._save(
            diffs=[("BOSS25-010", "Wire", 75.0), ("BOSS25-010", "Final Station HND", "Undercarriage")],
            additions=[{"SKU": "BOSS25-011", "Description": "NEW", "Wire": 61.0, "Trailer": ""}],
        )
        assert result == "abc1234"
        assert len(calls) == 1
        call = calls[0]
        assert call["path"] == "data/machine_clean.csv" and call["sha"] == "sha-fresh"
        assert call["message"] == "Update machine catalog via Excel import (2 cells, 1 new SKU(s))"
        out = pd.read_csv(io.StringIO(call["csv"])).set_index("SKU")
        assert out.at["BOSS25-010", "Wire"] == 75.0
        assert out.at["BOSS25-010", "Final Station HND"] == "Undercarriage"
        assert out.at["BOSS25-010", "Last Modified"] == app.today_local_str()
        # The concurrent edit on another row survives — the merge is cell-level.
        assert out.at["BOSS70-002", "Wire"] == 999.0
        assert pd.isna(out.at["BOSS70-002", "Last Modified"])
        assert out.at["PDS185-001", "Trailer"] == 0.0
        # The new row carries its values, blank where the sheet was blank.
        assert out.at["BOSS25-011", "Wire"] == 61.0 and out.at["BOSS25-011", "Description"] == "NEW"
        assert pd.isna(out.at["BOSS25-011", "Trailer"]) and pd.isna(out.at["BOSS25-011", "Warehouse"])
        assert list(out.index) == ["BOSS25-010", "BOSS70-002", "PDS185-001", "BOSS25-011"]

    def test_retries_once_on_a_sha_conflict(self, github):
        calls, state = github
        state["conflicts"] = 1
        assert self._save(diffs=[("BOSS25-010", "Wire", 75.0)], additions=[]) == "abc1234"
        assert len(calls) == 1 and calls[0]["message"] == "Update machine catalog via Excel import (1 cells)"

    def test_gives_up_after_two_conflicts(self, github):
        calls, state = github
        state["conflicts"] = 2
        assert self._save(diffs=[("BOSS25-010", "Wire", 75.0)], additions=[]) is None
        assert calls == []

    def test_nothing_to_save_makes_no_request(self, github):
        calls, _ = github
        assert self._save(diffs=[], additions=[]) is None
        assert calls == []

    def test_addition_already_on_file_is_skipped_not_overwritten(self, github):
        calls, _ = github
        result = self._save(diffs=[], additions=[{"SKU": "BOSS70-002", "Wire": 1.0}])
        # Nothing new to write once the fresh file already has it → no commit.
        assert result is None and calls == []

    def test_editor_path_still_computes_its_own_diff(self, github):
        calls, _ = github
        edited = _machine_live().reset_index(drop=True)
        edited.loc[edited["SKU"] == "PDS185-001", "Warehouse"] = 7
        app._save_catalog_csv(
            edited=edited, source_df=_machine_live(), editable_cols=list(excel_io.MACHINE_LABOR_COLUMNS),
            file_path="data/machine_clean.csv", label="machine", text_cols=["Description"],
        )
        assert calls[0]["message"] == "Update machine catalog via app (1 cells)"
        out = pd.read_csv(io.StringIO(calls[0]["csv"])).set_index("SKU")
        assert out.at["PDS185-001", "Warehouse"] == 7.0 and out.at["BOSS70-002", "Wire"] == 999.0


# ---------------------------------------------------------------------------
# The whole flow, against the real data files
# ---------------------------------------------------------------------------
class TestEndToEndRoundTrip:
    """What the feature exists for: export the database workbook, change a few
    labor minutes and add a SKU row *in that workbook*, upload it, see exactly
    those changes in the preview, save — and the CSV that goes to GitHub
    differs from the original in only those cells (plus the Last Modified
    stamps) and the one new row."""

    def test_export_edit_upload_save(self, monkeypatch):
        from openpyxl import load_workbook
        from core.data_loader import load_acc_labor, load_machine_labor

        exported = excel_io.build_database_workbook(include_saved_schedule_rows=False)
        wb = load_workbook(io.BytesIO(exported))
        ws = wb["Machine catalog"]
        header = [c.value for c in ws[1]]
        col = {name: i + 1 for i, name in enumerate(header)}

        def cell(r, name):
            return ws.cell(row=r, column=col[name])

        sku_a = cell(2, "SKU").value
        old_wire = float(cell(2, "Wire").value or 0)
        cell(2, "Wire").value = old_wire + 5              # typed as a number, as a planner would
        sku_b = cell(3, "SKU").value
        old_fn = float(cell(3, "FN_Assy").value or 0)
        cell(3, "FN_Assy").value = old_fn + 12.5
        assert cell(3, "Final Station CYP").value != "GenAcc"
        cell(3, "Final Station CYP").value = "genacc"      # wrong case on purpose
        new_row = [""] * len(header)
        for name, value in (("SKU", "ZZTEST-001"), ("Description", "ROUND TRIP TEST"),
                            ("Warehouse", 7), ("Bat", 2)):
            new_row[col[name] - 1] = value
        ws.append(new_row)
        buf = io.BytesIO()
        wb.save(buf)

        # Upload → reader (live columns come from data/) → preview.
        book = read_catalog_workbook(buf.getvalue(), "edited.xlsx")
        assert set(book.sheets) == {"machine", "accessory"}
        machine_sheet = book.sheets["machine"]
        assert machine_sheet.missing_columns == [] and machine_sheet.extra_columns == []
        machine_df = load_machine_labor()
        diff = diff_catalog_sheet(machine_sheet, machine_df)
        assert diff.cell_tuples() == [
            (sku_a, "Wire", old_wire + 5),
            (sku_b, "FN_Assy", old_fn + 12.5),
            (sku_b, "Final Station CYP", "GenAcc"),
        ]
        assert diff.new_skus == ["ZZTEST-001"]
        assert not diff.blocked and diff.absent_skus == []
        # The accessory sheet went through untouched: nothing to do, nothing wrong.
        acc = diff_catalog_sheet(book.sheets["accessory"], load_acc_labor())
        assert acc.is_empty and not acc.blocked and acc.absent_skus == []

        # Save, with GitHub stubbed so the "fresh" file is the real CSV.
        original = pd.read_csv(excel_io.DATA_DIR / "machine_clean.csv")
        calls = []

        def fake_save(csv_text, file_path, token, message, sha=None):
            calls.append((csv_text, message, sha))
            return {"commit": {"sha": "deadbeef0", "html_url": "u"}}

        monkeypatch.setattr(app, "_read_fresh_csv", lambda fp, tok: (original.copy(), "sha"))
        monkeypatch.setattr(app, "save_catalog_to_github", fake_save)
        monkeypatch.setattr(app, "_track_and_flush", lambda *a, **k: None)
        monkeypatch.setattr(app.st, "secrets", {"github_token": "t"})
        numeric, text = excel_io.editable_columns(machine_sheet.spec, machine_sheet.live_columns)
        result = app._save_catalog_csv(
            edited=None, source_df=machine_df, editable_cols=numeric,
            file_path="data/machine_clean.csv", label="machine", text_cols=text,
            diffs=diff.cell_tuples(), additions=diff.new_rows, via="Excel import",
        )
        assert result == "deadbee"
        (csv_text, message, sha), = calls
        assert sha == "sha"
        assert message == "Update machine catalog via Excel import (3 cells, 1 new SKU(s))"

        pushed = pd.read_csv(io.StringIO(csv_text))
        assert list(pushed.columns) == list(original.columns)
        assert len(pushed) == len(original) + 1
        before = original.set_index("SKU")
        after = pushed.iloc[:len(original)].set_index("SKU")
        assert list(after.index) == list(before.index)
        today = app.today_local_str()
        expected = {(sku_a, "Wire"), (sku_b, "FN_Assy"), (sku_b, "Final Station CYP")}
        for sku in (sku_a, sku_b):
            if str(before.at[sku, "Last Modified"]) != today:
                expected.add((sku, "Last Modified"))
        found = set()
        for column in before.columns:
            b, a = before[column], after[column]
            same = (b == a) | (b.isna() & a.isna())
            found.update((sku, column) for sku in before.index[~same.to_numpy()])
        assert found == expected
        assert after.at[sku_a, "Wire"] == old_wire + 5
        assert after.at[sku_b, "FN_Assy"] == old_fn + 12.5
        assert after.at[sku_b, "Final Station CYP"] == "GenAcc"
        assert after.at[sku_a, "Last Modified"] == today
        new = pushed.iloc[-1]
        assert new["SKU"] == "ZZTEST-001" and new["Description"] == "ROUND TRIP TEST"
        assert new["Warehouse"] == 7 and new["Bat"] == 2 and pd.isna(new["Wire"])
        assert new["Last Modified"] == today


class TestImportViewRenders:
    def test_import_sub_view_renders_without_an_upload(self, tmp_path, monkeypatch):
        """A full app run: switch 📁 Data & Setup to 📥 Import from Excel and
        make sure the view comes up (imports wired, nothing raises)."""
        AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
        import core.usage_tracker as usage_tracker
        monkeypatch.setattr(usage_tracker, "USAGE_LOG_PATH", tmp_path / "usage_log.jsonl")

        at = AppTest.from_file(str(APP), default_timeout=300)
        # Only the tab in front is built, so bring 📁 Data & Setup up first;
        # its view picker doesn't exist while 🏠 Overview is showing.
        at.session_state[app.ACTIVE_TAB_KEY] = app.TAB_DATA_SETUP
        at.run()
        assert not at.exception
        view = next(r for r in at.main.radio if "📥 Import from Excel" in list(r.options))
        view.set_value("📥 Import from Excel").run()
        assert not at.exception
        assert any("Import from Excel" in (s.value or "") for s in at.main.subheader)

    def test_preview_section_renders_for_an_edited_export(self, tmp_path, monkeypatch):
        """Render the post-upload preview with real widgets: metrics, the
        change table, the new-row table, the untouched-SKU expander and an
        enabled save button — for a workbook edited the way a planner would."""
        AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
        import core.usage_tracker as usage_tracker
        monkeypatch.setattr(usage_tracker, "USAGE_LOG_PATH", tmp_path / "usage_log.jsonl")
        # The stale-data banner polls GitHub; keep the test offline.
        monkeypatch.setattr(app, "_render_stale_data_banner", lambda *_a, **_k: None)

        from openpyxl import load_workbook
        exported = excel_io.build_database_workbook(include_saved_schedule_rows=False)
        wb = load_workbook(io.BytesIO(exported))
        ws = wb["Machine catalog"]
        col = {c.value: i + 1 for i, c in enumerate(ws[1])}
        ws.cell(2, col["Wire"]).value = float(ws.cell(2, col["Wire"]).value or 0) + 5
        ws.cell(2, col["Description"]).value = "EDITED IN EXCEL"
        ws.delete_rows(3)  # one SKU taken off the sheet — must be reported, not deleted
        new_row = [""] * len(col)
        new_row[col["SKU"] - 1] = "ZZTEST-002"
        new_row[col["Warehouse"] - 1] = 3
        ws.append(new_row)
        buf = io.BytesIO()
        wb.save(buf)
        edited = tmp_path / "edited.xlsx"
        edited.write_bytes(buf.getvalue())

        script = f"""
import app
from core import excel_io
from core.data_loader import load_machine_labor
book = excel_io.read_catalog_workbook(open({str(edited)!r}, "rb").read(), "edited.xlsx")
app._render_catalog_import_section(book.sheets["machine"], load_machine_labor(), "sig-1")
"""
        at = AppTest.from_string(script, default_timeout=300)
        at.run()
        assert not at.exception
        metrics = {m.label: m.value for m in at.metric}
        assert metrics["Cells to change"] == "2"
        assert metrics["New SKUs"] == "1"
        assert metrics["Catalog SKUs not on sheet"] == "1"
        labels = [b.label for b in at.button]
        assert any(l.startswith("💾 Save 2 cell change(s) + 1 new SKU(s)") for l in labels), labels
        assert all(not b.disabled for b in at.button)
        # Change table + new-row table rendered; no problems table.
        frames = [f.value for f in at.dataframe]
        assert any(list(f.columns) == ["SKU", "Column", "Old", "New", "Δ"] and len(f) == 2 for f in frames)
        assert any("SKU" in f.columns and list(f["SKU"]) == ["ZZTEST-002"] for f in frames)
        assert not any("Problem" in f.columns for f in frames)
        assert not at.error

    def test_preview_section_blocks_on_a_bad_cell(self, tmp_path, monkeypatch):
        AppTest = pytest.importorskip("streamlit.testing.v1").AppTest
        import core.usage_tracker as usage_tracker
        monkeypatch.setattr(usage_tracker, "USAGE_LOG_PATH", tmp_path / "usage_log.jsonl")
        monkeypatch.setattr(app, "_render_stale_data_banner", lambda *_a, **_k: None)

        from openpyxl import load_workbook
        exported = excel_io.build_database_workbook(include_saved_schedule_rows=False)
        wb = load_workbook(io.BytesIO(exported))
        ws = wb["Machine catalog"]
        col = {c.value: i + 1 for i, c in enumerate(ws[1])}
        ws.cell(2, col["Wire"]).value = "4O"      # letter O for zero
        ws.cell(4, col["QC"]).value = 99           # a good edit elsewhere — still held back
        buf = io.BytesIO()
        wb.save(buf)
        edited = tmp_path / "bad.xlsx"
        edited.write_bytes(buf.getvalue())

        script = f"""
import app
from core import excel_io
from core.data_loader import load_machine_labor
book = excel_io.read_catalog_workbook(open({str(edited)!r}, "rb").read(), "bad.xlsx")
app._render_catalog_import_section(book.sheets["machine"], load_machine_labor(), "sig-2")
"""
        at = AppTest.from_string(script, default_timeout=300)
        at.run()
        assert not at.exception
        assert any("didn't validate" in (e.value or "") for e in at.error)
        frames = [f.value for f in at.dataframe]
        problems = next(f for f in frames if "Problem" in f.columns)
        assert problems["Value"].tolist() == ["4O"]
        save = [b for b in at.button if b.label.startswith("💾 Save")]
        assert len(save) == 1 and save[0].disabled


# ---------------------------------------------------------------------------
# The IE workbook as it is — its own headers, not the app's export
# ---------------------------------------------------------------------------
IE_MACHINE_HEADER = [
    "🟡", "Unit", "SKU ID", "Description", "Warehouse (PICK)", "Wire Assy (PREP-WH)",
    "Sub Trl (SUB-TRL)", "Final Assy (FN-ASSY)", "PDI (TEST-PDI)", "QC (TEST-FN-QC)",
    "Ship (MAT-REIC-SHIP)", "Total", "🔋Bat", "Lead", "Flag",
]
IE_ACC_HEADER = [
    "SKU ID", "Description", "Warehouse (PICK)", "Acc KIT (PREP-ACC)",
    "Battery Assy (PREP-NP)", "Battery Sub raw (SUB-BTR raw)", "PM Acc (SUB-PM)",
    "Gen Sub (SUB-GEN)", "Compressor (SUB-COM)", "Total",
]


class TestIEWorkbookNativeLayout:
    """The industrial engineers' cycle-time workbook imports directly: a title
    row, their header spellings, marker / total / lead columns. Nothing has to
    be exported from the app first."""

    def _machine_book(self, sheet_name="Cycle Times"):
        rows = [
            ["Cycle Times — Machines"] + [None] * (len(IE_MACHINE_HEADER) - 1),
            [None] * len(IE_MACHINE_HEADER),
            IE_MACHINE_HEADER,
            # BOSS25-010: only FN_Assy differs from the live catalog (45 → 50).
            [None, "EA", "BOSS25-010", "EBOSS 25 TRAILER", 10, 60, 210, 50, 60, 60, 30, 480, 1, None, None],
            # A SKU the catalog doesn't have.
            [None, "EA", "NEW-001", "brand new", 5, 0, 0, 100, 30, 30, 15, 180, 0, None, None],
        ]
        return _workbook({sheet_name: rows})

    def test_unnamed_sheet_is_detected_by_its_ie_header(self):
        book = read_catalog_workbook(self._machine_book(), "CycleTimes.xlsx", live_columns=LIVE)
        assert list(book.sheets) == ["machine"]
        assert book.sheets["machine"].matched_by == "header"

    def test_every_ie_header_lands_on_the_right_column_and_is_reported(self):
        sheet = _sheet(self._machine_book())
        assert sheet.aliased == {
            "SKU ID": "SKU",
            "Warehouse (PICK)": "Warehouse",
            "Wire Assy (PREP-WH)": "Wire",
            "Sub Trl (SUB-TRL)": "Trailer",
            "Final Assy (FN-ASSY)": "FN_Assy",
            "PDI (TEST-PDI)": "PDI",
            "QC (TEST-FN-QC)": "QC",
            "Ship (MAT-REIC-SHIP)": "Ship",
            "🔋Bat": "Bat",
        }
        # "Description" is the app's own name — matched exactly, not aliased.
        assert "Description" in sheet.columns and "Description" not in sheet.aliased
        assert set(sheet.extra_columns) == {"🟡", "Unit", "Total", "Lead", "Flag"}
        # The IE sheet has no ETO and no routing — those stay untouched.
        assert "ETO" in sheet.missing_columns
        assert "Final Station HND" in sheet.missing_columns
        assert "Warehouse" not in sheet.missing_columns

    def test_diff_changes_exactly_the_edited_minute_and_adds_the_new_sku(self):
        diff = diff_catalog_sheet(_sheet(self._machine_book()), _machine_live())
        assert not diff.blocked
        assert diff.cell_tuples() == [("BOSS25-010", "FN_Assy", 50.0)]
        assert diff.new_skus == ["NEW-001"]
        # Catalog SKUs the IE sheet doesn't list are left alone — placeholders survive.
        assert set(diff.absent_skus) == {"BOSS70-002", "PDS185-001"}

    def test_named_ie_sheet_also_reconciles_through_the_aliases(self):
        # "Machine" is a recognised sheet name, so this one is matched by name —
        # the header aliasing has to kick in regardless of how the sheet was found.
        sheet = _sheet(self._machine_book(sheet_name="Machine"))
        assert sheet.matched_by == "name"
        assert sheet.aliased["Final Assy (FN-ASSY)"] == "FN_Assy"
        assert diff_catalog_sheet(sheet, _machine_live()).cell_tuples() == [
            ("BOSS25-010", "FN_Assy", 50.0),
        ]

    def test_accessory_sheet_in_ie_spelling(self):
        rows = [
            IE_ACC_HEADER,
            # GenAcc 0 → 145; everything else equal to the live row.
            ["BOSS25-A016", "KIT A", 10, 10, 0, 320, 60, 145, 0, 545],
        ]
        book = read_catalog_workbook(_workbook({"Acc": rows}), "CycleTimes.xlsx", live_columns=LIVE)
        sheet = book.sheets["accessory"]
        assert sheet.aliased == {
            "SKU ID": "SKU",
            "Warehouse (PICK)": "Warehouse",
            "Acc KIT (PREP-ACC)": "AccKIT",
            "Battery Assy (PREP-NP)": "Nameplate Prep",
            "Battery Sub raw (SUB-BTR raw)": "BattSubRaw",
            "PM Acc (SUB-PM)": "PMAcc",
            "Gen Sub (SUB-GEN)": "GenAcc",
            "Compressor (SUB-COM)": "ComAcc",
        }
        assert diff_catalog_sheet(sheet, _acc_live()).cell_tuples() == [
            ("BOSS25-A016", "GenAcc", 145.0),
        ]

    def test_a_schedule_tab_is_not_mistaken_for_a_catalog(self):
        # A schedule sheet has "SKU" inside other headers ("SKU IN BC",
        # "Cycle Time FG SKU ID") but no SKU header cell of its own. The SKU
        # header test is exact on purpose so this tab is skipped, not imported.
        rows = [
            ["PRIORITY", "LOCATION", "MODEL TYPE", "FG SKU ID", "BUILD QTY",
             "SKU IN BC", "Cycle Time FG SKU ID", "Total Time"],
            [1, "Spartanburg", "SDG45S", "SDG45-001", 1, "YES", 240, 300],
        ]
        book = read_catalog_workbook(_workbook({"Schedule": rows}), "sched.xlsx", live_columns=LIVE)
        assert book.sheets == {}
        assert book.skipped and "SKU" in book.skipped[0][1]

    def test_routing_header_is_never_aliased_onto_final_assembly(self):
        # With an empty live header the fallback live columns have no routing
        # names, so "Final Station HND" can't match exactly — it must then be
        # ignored, not aliased to FN_Assy through "Final".
        rows = [["SKU", "FN_Assy", "Final Station HND"], ["BOSS25-010", 45, "Undercarriage"]]
        book = read_catalog_workbook(
            _workbook({"Machine catalog": rows}), "x.xlsx",
            live_columns={"machine": [], "accessory": []},
        )
        sheet = book.sheets["machine"]
        assert "Final Station HND" in sheet.extra_columns
        assert sheet.aliased == {}


# ---------------------------------------------------------------------------
# Reconcile: temporary in-app minute edits vs what the sheet says
# ---------------------------------------------------------------------------
from core.excel_io import reconcile_pending  # noqa: E402


class TestReconcilePending:
    """Excel is the source of truth, but a planner may override a minute in
    the app and update the workbook later. The catalog holds the override, so
    a sheet that still has the old number would — without this — silently
    revert it on import."""

    def _live_with_override(self):
        # The planner set BOSS25-010 FN_Assy 45 → 150 in the app; Excel had 45.
        live = _machine_live()
        live.loc["BOSS25-010", "FN_Assy"] = 150.0
        return live

    PENDING = {"BOSS25-010": {"FN_Assy": {"app_value": 150.0, "excel_value": 45.0, "at": "t"}}}

    def _run(self, fn_assy_on_sheet, pending=None, revert=False, include_sku=True):
        rows = _machine_rows(
            *([["BOSS25-010", "EBOSS 25 TRAILER", 10, 60, 210, fn_assy_on_sheet, 60, 60, 30, 0, 1]]
              if include_sku else []),
            ["BOSS70-002", "EBOSS 70", 10, 90, 240, 0, 60, 60, 30, 0, 3],
        )
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        diff = diff_catalog_sheet(sheet, self._live_with_override())
        return reconcile_pending(diff, sheet, self.PENDING if pending is None else pending,
                                 revert_to_excel=revert)

    def test_excel_still_old_keeps_the_app_value_and_drops_the_reverting_cell(self):
        diff, report = self._run(fn_assy_on_sheet=45)
        assert report.still_pending == [("BOSS25-010", "FN_Assy", 150.0, 45.0)]
        assert diff.cell_tuples() == []           # the 150 → 45 revert is gone
        assert report.to_clear == []              # still pending in Excel

    def test_excel_caught_up_clears_quietly(self):
        diff, report = self._run(fn_assy_on_sheet=150)
        assert report.caught_up == [("BOSS25-010", "FN_Assy", 150.0, 150.0)]
        assert diff.cell_tuples() == []           # nothing differs
        assert report.to_clear == [("BOSS25-010", "FN_Assy")]

    def test_excel_changed_to_something_else_wins(self):
        diff, report = self._run(fn_assy_on_sheet=160)
        assert report.excel_changed == [("BOSS25-010", "FN_Assy", 150.0, 160.0)]
        assert diff.cell_tuples() == [("BOSS25-010", "FN_Assy", 160.0)]
        assert report.to_clear == [("BOSS25-010", "FN_Assy")]

    def test_revert_applies_excels_old_value_and_clears(self):
        diff, report = self._run(fn_assy_on_sheet=45, revert=True)
        assert report.reverted == [("BOSS25-010", "FN_Assy", 150.0, 45.0)]
        assert diff.cell_tuples() == [("BOSS25-010", "FN_Assy", 45.0)]
        assert report.to_clear == [("BOSS25-010", "FN_Assy")]

    def test_sku_not_on_sheet_is_untouched(self):
        diff, report = self._run(fn_assy_on_sheet=45, include_sku=False)
        assert report.not_on_sheet == [("BOSS25-010", "FN_Assy", 150.0)]
        assert report.to_clear == []

    def test_blank_cell_on_sheet_when_excel_had_a_value_is_untouched(self):
        diff, report = self._run(fn_assy_on_sheet=None)
        assert report.not_on_sheet == [("BOSS25-010", "FN_Assy", 150.0)]
        assert diff.cell_tuples() == []

    def test_blank_cell_when_excel_was_blank_is_still_pending(self):
        pending = {"BOSS25-010": {"FN_Assy": {"app_value": 150.0, "excel_value": None, "at": "t"}}}
        _diff, report = self._run(fn_assy_on_sheet=None, pending=pending)
        assert report.still_pending == [("BOSS25-010", "FN_Assy", 150.0, None)]

    def test_other_cells_on_the_same_row_still_import(self):
        # Only the pending cell is protected; a genuine change elsewhere lands.
        rows = _machine_rows(["BOSS25-010", "EBOSS 25 TRAILER", 10, 60, 210, 45, 90, 60, 30, 0, 1])
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        diff = diff_catalog_sheet(sheet, self._live_with_override())
        diff, report = reconcile_pending(diff, sheet, self.PENDING)
        assert diff.cell_tuples() == [("BOSS25-010", "PDI", 90.0)]
        assert report.still_pending and not report.caught_up

    def test_nothing_pending_is_a_passthrough(self):
        rows = _machine_rows(["BOSS25-010", "EBOSS 25 TRAILER", 10, 60, 210, 45, 60, 60, 30, 0, 1])
        sheet = _sheet(_workbook({"Machine catalog": rows}))
        diff = diff_catalog_sheet(sheet, self._live_with_override())
        same, report = reconcile_pending(diff, sheet, {})
        assert same is diff and report.is_empty

    def test_sku_match_is_case_insensitive(self):
        pending = {"boss25-010": {"FN_Assy": {"app_value": 150.0, "excel_value": 45.0, "at": "t"}}}
        diff, report = self._run(fn_assy_on_sheet=45, pending=pending)
        assert report.still_pending and diff.cell_tuples() == []
