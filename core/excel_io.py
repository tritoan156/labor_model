"""Excel bridge — read planner workbooks in, write the whole database out.

Two jobs, both Streamlit-free so they're unit-testable:

1. **Inbound** — :func:`upload_to_schedule_csv` turns an uploaded ``.xlsx`` /
   ``.xlsm`` workbook into the CSV bytes the schedule loader already
   understands. It scans every sheet for the first header row that carries
   the five required schedule columns (so a title row above the header, or
   the schedule living on the third tab, both just work) and normalizes Excel
   cell types (dates → ``YYYY-MM-DD``, ``5.0`` → ``5``) on the way through.
   A ``.csv`` passes straight through untouched, so every downstream path —
   the session's active schedule, the 💾 Save schedule button, the saved
   schedule JSON — keeps seeing CSV and nothing else has to change.

2. **Outbound** — :func:`database_tables` flattens every file under ``data/``
   (the four CSV catalogs plus the crew / settings / scenarios / executive
   scenarios / saved schedules / process-flow JSON) into plain DataFrames, and
   :func:`build_workbook` writes any ``{sheet: DataFrame}`` map to a single
   ``.xlsx`` with bold frozen headers and autofilters. The app uses it twice:
   once for the whole database, once for the current session's analysis
   (schedule / per-unit labor / capacity table).

:func:`raw_github_url` and :func:`power_query_m_for_csv` back the "link Excel
to the live data" panel — the data files are plain CSV/JSON on GitHub, so
Excel's *Get Data → From Web* can read them directly and refresh on demand.
"""
from __future__ import annotations

import io
import json
import math
import re
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd

from .catalog_storage import GITHUB_BRANCH, GITHUB_REPO
from .data_loader import (
    ScheduleColumnError,
    _REQUIRED_SCHEDULE_COLUMNS,
    _find_column,
)
from .exec_scenario_storage import load_all_exec_scenarios
from .facility_settings_storage import load_all_facility_settings
from .facility_storage import load_all_facility_configs
from .process_flow_storage import load_process_flow
from .scenario_storage import load_all_scenarios
from .uploaded_schedule_storage import load_all_uploaded_schedules

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

#: Workbook extensions the uploader accepts. ``.xls`` (the pre-2007 binary
#: format) needs the ``xlrd`` package, which isn't a dependency — planners
#: with one of those save it as ``.xlsx`` first.
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")

#: How many leading rows of a sheet to scan for the header row. Real-world
#: exports put a title, a date, and a blank line above the column names; a
#: header deeper than this is almost certainly not a schedule sheet.
HEADER_SCAN_ROWS = 20

#: Excel caps sheet names at 31 characters and bans these characters.
_SHEET_NAME_BAD_CHARS = re.compile(r"[\[\]:*?/\\]")
_SHEET_NAME_MAX = 31


# ---------------------------------------------------------------------------
# Inbound — workbook → schedule CSV bytes
# ---------------------------------------------------------------------------
def is_excel_filename(name) -> bool:
    """True when ``name`` ends in an extension :func:`upload_to_schedule_csv`
    will open as a workbook (case-insensitive)."""
    return str(name or "").strip().lower().endswith(EXCEL_EXTENSIONS)


def _cell_to_text(value) -> str:
    """Render one Excel cell the way a CSV export would.

    * blanks / NaN / NaT → ``""``
    * dates and datetimes → ``YYYY-MM-DD`` (the loader parses this for both
      PRODUCTION MONTH and PRODUCTION DAY; the midnight time part a date cell
      carries is noise)
    * whole-number floats → ``"5"`` not ``"5.0"`` (BUILD QTY reads cleaner in
      the saved CSV; the loader coerces either)
    * everything else → ``str(value)``, stripped
    """
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass  # non-scalar (shouldn't happen for a cell) — fall through
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value).strip()


def _row_matches_schedule_header(cells: list) -> bool:
    """True when every required schedule column fuzzy-matches one of ``cells``."""
    names = [_cell_to_text(c) for c in cells]
    names = [n for n in names if n]
    if not names:
        return False
    return all(
        _find_column(names, *patterns) is not None
        for _canonical, patterns, _label in _REQUIRED_SCHEDULE_COLUMNS
    )


def find_schedule_header_row(raw: pd.DataFrame) -> int | None:
    """Index of the first row (within :data:`HEADER_SCAN_ROWS`) that looks like
    a schedule header, or ``None`` if the sheet has no such row.

    ``raw`` is a sheet read with ``header=None`` so row 0 is the first
    spreadsheet row.
    """
    if raw is None or raw.empty:
        return None
    limit = min(len(raw), HEADER_SCAN_ROWS)
    for i in range(limit):
        if _row_matches_schedule_header(raw.iloc[i].tolist()):
            return i
    return None


def _sheet_to_csv_bytes(raw: pd.DataFrame, header_row: int) -> bytes:
    """Re-head ``raw`` at ``header_row``, drop empty rows, stringify cells, and
    serialize as UTF-8 CSV."""
    header = [_cell_to_text(c) for c in raw.iloc[header_row].tolist()]
    # Blank header cells get a stable placeholder so two of them don't collide
    # into one duplicate "" column (pandas would do the same on read_csv).
    header = [h if h else f"Unnamed: {i}" for i, h in enumerate(header)]
    body = raw.iloc[header_row + 1:].copy()
    body.columns = header
    body = body.dropna(how="all")
    for col in list(body.columns):
        body[col] = body[col].map(_cell_to_text)
    # A row that's blank in every cell after stringifying is padding — drop it.
    body = body[(body != "").any(axis=1)]
    return body.to_csv(index=False).encode("utf-8")


def excel_to_schedule_csv(raw_bytes: bytes, filename: str = "") -> tuple[bytes, str]:
    """Convert an ``.xlsx`` / ``.xlsm`` workbook to schedule CSV bytes.

    Returns ``(csv_bytes, sheet_name)`` for the first sheet, in workbook order,
    that has a header row carrying all five required schedule columns.

    Raises :class:`ScheduleColumnError` (the same planner-friendly error the
    CSV loader raises, so the upload UI surfaces it the same way) when the
    file can't be opened as a workbook or no sheet has a usable header.
    """
    label = f"**{filename}**" if filename else "the workbook"
    try:
        # ExcelFile as a context manager so the underlying zip is closed
        # deterministically (a bare read_excel on a BytesIO leaves it to the
        # GC and trips a ResourceWarning on Python 3.13).
        with pd.ExcelFile(io.BytesIO(raw_bytes), engine="openpyxl") as xf:
            book = {
                name: xf.parse(name, header=None, dtype=object)
                for name in xf.sheet_names
            }
    except Exception as exc:  # noqa: BLE001 — surface whatever openpyxl said
        raise ScheduleColumnError(
            f"We couldn't open {label} as an Excel workbook ({exc}). "
            "Save it as .xlsx (or export the sheet as CSV) and re-upload."
        ) from exc
    if not book:
        raise ScheduleColumnError(f"{label} has no sheets.")

    skipped: list[str] = []
    for sheet_name, raw in book.items():
        header_row = find_schedule_header_row(raw)
        if header_row is None:
            skipped.append(str(sheet_name))
            continue
        return _sheet_to_csv_bytes(raw, header_row), str(sheet_name)

    tried = ", ".join(f"`{s}`" for s in skipped) or "—"
    raise ScheduleColumnError(
        f"None of the sheets in {label} ({tried}) has a header row with the "
        "columns we need: **Location, FG SKU ID, Accessory SKU ID, Build Qty, "
        "Production Month** (any of their usual spellings). Put those column "
        "names in a single row on the schedule sheet and re-upload."
    )


def upload_to_schedule_csv(filename: str, raw_bytes: bytes) -> tuple[bytes, str | None]:
    """Normalize an uploaded schedule file to CSV bytes.

    * Excel workbook → converted via :func:`excel_to_schedule_csv`; returns
      ``(csv_bytes, sheet_name)``.
    * Anything else (CSV) → passed through untouched; returns
      ``(raw_bytes, None)``.
    """
    if is_excel_filename(filename):
        return excel_to_schedule_csv(raw_bytes, filename)
    return raw_bytes, None


# ---------------------------------------------------------------------------
# Outbound — every data file as a flat table
# ---------------------------------------------------------------------------
#: (sheet name, source file under data/, one-line description). Order is the
#: sheet order in the exported workbook and the row order in the link table.
DATABASE_TABLES: list[tuple[str, str, str]] = [
    ("Machine catalog", "machine_clean.csv",
     "FG SKU labor minutes by station, battery count, per-facility routing"),
    ("Accessory catalog", "acc_clean.csv",
     "Accessory SKU labor minutes by station"),
    ("Items", "item_master.csv",
     "Item master — defaults, family variants, per-accessory entries"),
    ("Item packages", "item_packages.csv",
     "Named item bundles (CWP / NE / AWP) and their components"),
    ("Crew", "facility_crew.json",
     "Headcount / concurrency / crew per station, one row per facility × station"),
    ("Facility settings", "facility_settings.json",
     "Per-facility planning assumptions (new-hire share, absenteeism, efficiency…)"),
    ("Scenarios", "scenarios.json",
     "Saved manual build plans — one row per SKU line"),
    ("Exec scenarios", "exec_scenarios.json",
     "Saved executive summaries — assumptions and frozen metrics, one row per value"),
    ("Saved schedules", "uploaded_schedules.json",
     "Index of saved uploaded schedules (facility, name, saved at, rows)"),
    ("Saved schedule rows", "uploaded_schedules.json",
     "Every row of every saved uploaded schedule"),
    ("Process flow", "process_flow.json",
     "Station precedence edges by unit class"),
    ("Sample schedule", "may_schedule.csv",
     "Bundled sample schedule (scrubbed)"),
]

#: Files Excel can link to with one *From Web* step (plain CSV). The JSON
#: files need a few Power Query transforms — the exported workbook already
#: has them flattened, so the link table points planners there instead.
CSV_LINKABLE_FILES = (
    "machine_clean.csv", "acc_clean.csv", "item_master.csv",
    "item_packages.csv", "may_schedule.csv",
)


def raw_github_url(file_name: str, repo: str = GITHUB_REPO,
                   branch: str = GITHUB_BRANCH) -> str:
    """Raw-content URL for ``data/<file_name>`` on GitHub — what Excel's
    *From Web* connector reads."""
    return f"https://raw.githubusercontent.com/{repo}/{branch}/data/{file_name}"


def power_query_m_for_csv(url: str) -> str:
    """A ready-to-paste Power Query (M) query that loads a UTF-8 CSV from
    ``url`` and promotes the first row to headers."""
    return (
        "let\n"
        f'    Source = Csv.Document(Web.Contents("{url}"),\n'
        "        [Delimiter=\",\", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),\n"
        "    Headers = Table.PromoteHeaders(Source, [PromoteAllScalars=true])\n"
        "in\n"
        "    Headers"
    )


def _read_csv_raw(path: Path) -> pd.DataFrame:
    """Read a catalog CSV exactly as stored (all text, blanks kept) so the
    export round-trips the file byte-for-byte in spirit — no loader-side
    renames, fills or derived columns."""
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")


def crew_table(cfg: dict | None) -> pd.DataFrame:
    """``facility_crew.json`` → Facility / Station / HC / Conc / Crew rows."""
    rows = []
    for facility, stations in (cfg or {}).items():
        for station, vals in (stations or {}).items():
            vals = vals or {}
            rows.append({
                "Facility": facility, "Station": station,
                "HC": vals.get("HC"), "Conc": vals.get("Conc"), "Crew": vals.get("Crew"),
            })
    return pd.DataFrame(rows, columns=["Facility", "Station", "HC", "Conc", "Crew"])


def settings_table(settings: dict | None) -> pd.DataFrame:
    """``facility_settings.json`` → Facility / Setting / Value rows (lists are
    joined with ``; `` so they stay one cell)."""
    rows = []
    for facility, vals in (settings or {}).items():
        for key, val in (vals or {}).items():
            if isinstance(val, (list, tuple, set)):
                val = "; ".join(str(v) for v in val)
            elif isinstance(val, dict):
                val = json.dumps(val)
            rows.append({"Facility": facility, "Setting": key, "Value": val})
    return pd.DataFrame(rows, columns=["Facility", "Setting", "Value"])


def scenarios_table(scenarios: dict | None) -> pd.DataFrame:
    """``scenarios.json`` → one row per saved SKU line, tagged with facility,
    scenario name and save timestamp."""
    rows = []
    for facility, by_name in (scenarios or {}).items():
        for name, block in (by_name or {}).items():
            block = block or {}
            saved_at = block.get("saved_at", "")
            entries = block.get("entries") or []
            if not entries:
                rows.append({"Facility": facility, "Scenario": name, "Saved at": saved_at})
                continue
            for entry in entries:
                row = {"Facility": facility, "Scenario": name, "Saved at": saved_at}
                row.update(entry or {})
                rows.append(row)
    df = pd.DataFrame(rows)
    lead = [c for c in ("Facility", "Scenario", "Saved at") if c in df.columns]
    rest = [c for c in df.columns if c not in lead]
    return df[lead + rest] if not df.empty else pd.DataFrame(
        columns=["Facility", "Scenario", "Saved at", "FG SKU", "Accessory SKU", "Qty"])


def exec_scenarios_table(exec_scenarios: dict | None) -> pd.DataFrame:
    """``exec_scenarios.json`` → long format: one row per assumption or
    metric value. Long beats wide here — the metric set has grown over time,
    and a key/value layout pivots cleanly in Excel without blank columns."""
    rows = []
    for facility, by_name in (exec_scenarios or {}).items():
        for name, block in (by_name or {}).items():
            block = block or {}
            saved_at = block.get("saved_at", "")
            for section_key, section_label in (("assumptions", "Assumption"),
                                                ("metrics", "Metric")):
                for key, val in (block.get(section_key) or {}).items():
                    if isinstance(val, (dict, list, tuple, set)):
                        val = json.dumps(val, default=str)
                    rows.append({
                        "Facility": facility, "Scenario": name, "Saved at": saved_at,
                        "Section": section_label, "Name": key, "Value": val,
                    })
    return pd.DataFrame(
        rows, columns=["Facility", "Scenario", "Saved at", "Section", "Name", "Value"])


def saved_schedules_index(uploaded: dict | None) -> pd.DataFrame:
    """``uploaded_schedules.json`` → one row per saved schedule."""
    rows = []
    for facility, by_name in (uploaded or {}).items():
        for name, block in (by_name or {}).items():
            block = block or {}
            rows.append({
                "Facility": facility, "Schedule": name,
                "Saved at": block.get("saved_at", ""), "Rows": block.get("rows"),
            })
    return pd.DataFrame(rows, columns=["Facility", "Schedule", "Saved at", "Rows"])


def saved_schedule_rows(uploaded: dict | None) -> pd.DataFrame:
    """``uploaded_schedules.json`` → every saved schedule's rows stacked, with
    Facility + Schedule prepended. Columns are the union across schedules;
    all values stay text exactly as the planner's CSV carried them."""
    frames = []
    for facility, by_name in (uploaded or {}).items():
        for name, block in (by_name or {}).items():
            csv_text = (block or {}).get("csv_content") or ""
            if not csv_text.strip():
                continue
            try:
                df = pd.read_csv(io.StringIO(csv_text), dtype=str, keep_default_na=False)
            except Exception:  # noqa: BLE001 — one bad blob shouldn't sink the export
                continue
            df.columns = [str(c).lstrip("﻿").strip() for c in df.columns]
            df.insert(0, "Schedule", name)
            df.insert(0, "Facility", facility)
            frames.append(df)
    if not frames:
        return pd.DataFrame(columns=["Facility", "Schedule"])
    return pd.concat(frames, ignore_index=True, sort=False)


def process_flow_table(edges: list | None) -> pd.DataFrame:
    """``process_flow.json`` edges → From / To / Class rows."""
    df = pd.DataFrame(edges or [])
    if df.empty:
        return pd.DataFrame(columns=["From", "To", "Class"])
    lead = [c for c in ("From", "To", "Class") if c in df.columns]
    return df[lead + [c for c in df.columns if c not in lead]]


#: The one sheet that dominates export time and file size: every saved
#: schedule's rows (thousands of rows × every ERP column the planner's CSV
#: carried). Callers can leave it out for a fast "catalogs + config" export.
SAVED_SCHEDULE_ROWS_SHEET = "Saved schedule rows"


def database_tables(data_dir: Path | str | None = None, *,
                    include_saved_schedule_rows: bool = True) -> dict[str, pd.DataFrame]:
    """Every data file as a flat DataFrame, keyed by sheet name in
    :data:`DATABASE_TABLES` order.

    ``data_dir`` overrides the repo's ``data/`` for the CSVs (tests); the JSON
    loaders always read the repo copy, which is the same thing in production.
    ``include_saved_schedule_rows=False`` skips :data:`SAVED_SCHEDULE_ROWS_SHEET`
    (the index of saved schedules is always included).
    """
    data_dir = Path(data_dir) if data_dir is not None else DATA_DIR
    crew = load_all_facility_configs()
    settings = load_all_facility_settings()
    scenarios = load_all_scenarios()
    exec_scen = load_all_exec_scenarios()
    uploaded = load_all_uploaded_schedules()
    edges = load_process_flow()
    builders = {
        "Machine catalog": lambda: _read_csv_raw(data_dir / "machine_clean.csv"),
        "Accessory catalog": lambda: _read_csv_raw(data_dir / "acc_clean.csv"),
        "Items": lambda: _read_csv_raw(data_dir / "item_master.csv"),
        "Item packages": lambda: _read_csv_raw(data_dir / "item_packages.csv"),
        "Crew": lambda: crew_table(crew),
        "Facility settings": lambda: settings_table(settings),
        "Scenarios": lambda: scenarios_table(scenarios),
        "Exec scenarios": lambda: exec_scenarios_table(exec_scen),
        "Saved schedules": lambda: saved_schedules_index(uploaded),
        "Saved schedule rows": lambda: saved_schedule_rows(uploaded),
        "Process flow": lambda: process_flow_table(edges),
        "Sample schedule": lambda: _read_csv_raw(data_dir / "may_schedule.csv"),
    }
    return {
        sheet: builders[sheet]()
        for sheet, _file, _desc in DATABASE_TABLES
        if include_saved_schedule_rows or sheet != SAVED_SCHEDULE_ROWS_SHEET
    }


def database_readme(branch: str = GITHUB_BRANCH, *,
                    include_saved_schedule_rows: bool = True) -> pd.DataFrame:
    """The README sheet for the database export: what each sheet is, which
    file it came from, and the raw URL Excel can link to."""
    rows = []
    for sheet, file_name, desc in DATABASE_TABLES:
        if not include_saved_schedule_rows and sheet == SAVED_SCHEDULE_ROWS_SHEET:
            continue
        rows.append({
            "Sheet": sheet,
            "Source file": f"data/{file_name}",
            "What it holds": desc,
            "Live link (Excel → Data → From Web)":
                raw_github_url(file_name, branch=branch)
                if file_name in CSV_LINKABLE_FILES else "— JSON; use this sheet —",
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Workbook writer
# ---------------------------------------------------------------------------
def safe_sheet_name(name: str, taken: set | None = None) -> str:
    """Excel-legal sheet name: banned characters stripped, 31-char cap, and
    a numeric suffix if ``taken`` already holds it."""
    base = _SHEET_NAME_BAD_CHARS.sub("", str(name or "Sheet")).strip() or "Sheet"
    base = base[:_SHEET_NAME_MAX]
    if taken is None:
        return base
    candidate, n = base, 2
    while candidate in taken:
        suffix = f" ({n})"
        candidate = base[:_SHEET_NAME_MAX - len(suffix)] + suffix
        n += 1
    return candidate


def _excel_safe_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Make a frame openpyxl can write: tz-aware datetimes → naive, non-scalar
    objects (lists/dicts) → JSON text, index folded in when it's named."""
    if df is None:
        return pd.DataFrame()
    out = df.copy()
    if out.index.name is not None or isinstance(out.index, pd.MultiIndex):
        # Fold a named index in as leading columns; rename any level whose
        # name collides with an existing column so reset_index can't raise.
        names = [n if n is not None else f"level_{i}" for i, n in enumerate(out.index.names)]
        names = [f"{n} (index)" if n in out.columns else n for n in names]
        out.index = out.index.set_names(names)
        out = out.reset_index()
    # Duplicate column labels break both the per-column pass below (``out[c]``
    # would be a frame, not a series) and to_excel's header write — so
    # disambiguate first.
    if out.columns.duplicated().any():
        seen: dict = {}
        cols = []
        for c in out.columns:
            n = seen.get(c, 0)
            cols.append(c if n == 0 else f"{c}.{n}")
            seen[c] = n + 1
        out.columns = cols
    for col in out.columns:
        s = out[col]
        if isinstance(s.dtype, pd.DatetimeTZDtype):
            out[col] = s.dt.tz_localize(None)
        elif s.dtype == object:
            out[col] = s.map(
                lambda v: json.dumps(v, default=str)
                if isinstance(v, (list, dict, set, tuple)) else v
            )
    return out


def _polish_sheet(ws, n_cols: int, n_rows: int) -> None:
    """Bold header, freeze it, add an autofilter, and size columns to content
    (capped so a long description doesn't produce a screen-wide column)."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if n_cols == 0:
        return
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    if n_rows > 0:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"
    for idx in range(1, n_cols + 1):
        letter = get_column_letter(idx)
        longest = 0
        # Sample the header + first 200 rows; enough to size sensibly and
        # cheap enough for the 10k-row saved-schedule sheet.
        for row in ws.iter_rows(min_col=idx, max_col=idx, min_row=1,
                                max_row=min(n_rows + 1, 201), values_only=True):
            v = row[0]
            if v is not None:
                longest = max(longest, len(str(v)))
        ws.column_dimensions[letter].width = max(8, min(60, longest + 2))


def build_workbook(tables: dict[str, pd.DataFrame],
                   readme: pd.DataFrame | None = None) -> bytes:
    """Write ``{sheet name: DataFrame}`` to one ``.xlsx`` and return its bytes.

    ``readme`` (optional) becomes the first sheet, named ``README``. Sheet
    names are made Excel-legal and de-duplicated; an empty frame still gets a
    sheet (with just its headers) so the workbook layout is stable.
    """
    buf = io.BytesIO()
    taken: set = set()
    ordered: list[tuple[str, pd.DataFrame]] = []
    if readme is not None:
        ordered.append(("README", readme))
        taken.add("README")
    for name, df in tables.items():
        sheet = safe_sheet_name(name, taken)
        taken.add(sheet)
        ordered.append((sheet, df))

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for sheet, df in ordered:
            frame = _excel_safe_frame(df)
            frame.to_excel(writer, sheet_name=sheet, index=False)
            _polish_sheet(writer.sheets[sheet], len(frame.columns), len(frame))
    return buf.getvalue()


def build_database_workbook(data_dir: Path | str | None = None, *,
                            include_saved_schedule_rows: bool = True) -> bytes:
    """The whole ``data/`` folder as one workbook: README + one sheet per
    entry in :data:`DATABASE_TABLES` (the saved-schedule rows sheet only when
    ``include_saved_schedule_rows`` — it's the slow, big one)."""
    return build_workbook(
        database_tables(data_dir, include_saved_schedule_rows=include_saved_schedule_rows),
        readme=database_readme(include_saved_schedule_rows=include_saved_schedule_rows),
    )
